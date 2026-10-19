"""Fonte estadual da base Brasil a partir da planilha de médicos únicos por estado (uso administrativo).

A planilha entregue pelo responsável pelo produto em 19/10/2026 (medicos_unicos_municipios_estados_201401_202406.xlsx)
traz, na aba "Estados", uma tabela larga: uma linha por competência (AAAAMM) e uma coluna por UF (sigla), com a contagem
distinta de médicos dentro da UF e da competência, "NÃO soma de municípios". Este script grava a mesma informação, sem
alterar valores, em Dados_Estados_Brasil.csv, no formato das outras fontes da base: uma linha por UF e competência
(uf;periodo_ano;periodo_mes;categoria_profissional;contagem), separador ponto e vírgula e fim de linha CRLF, como a fonte
municipal (assim o SHA-256 não muda num checkout com conversão de fim de linha).

A conversão é recusada, sem gravar nada, quando: falta a aba "Estados"; o título ou o subtítulo da aba mudam (são eles
que dizem que a contagem é de médicos únicos por UF); o cabeçalho não é "Competência" seguido das 27 UFs, cada uma uma
vez; uma competência não é AAAAMM válida ou se repete; uma contagem está vazia, não é inteira ou não é positiva; há
célula preenchida fora da tabela além das notas da coluna A. A categoria MEDICINA vem do título; nada é estimado,
somado ou corrigido.

Uso, na raiz do repositório:
    python backend/build_state_source.py <planilha.xlsx>            # confere e mostra o resumo, sem gravar
    python backend/build_state_source.py <planilha.xlsx> --gravar   # grava Dados_Estados_Brasil.csv na raiz
Depois de gravar: python backend/brazil_data.py manifesto --gravar, revisão do diff e registro (AGENTS.md, regra 2).
"""
import argparse
import hashlib
import io
import json
import re
from pathlib import Path

from openpyxl import load_workbook

from brazil_data import FILES, REPO_ROOT
from territorial_selection import STATES

SHEET = "Estados"
TITLE = "MÉDICOS ÚNICOS POR ESTADO"
SUBTITLE = "Contagem distinta de chaves dentro da UF e competência — NÃO soma de municípios"
HEADER = "Competência"
CATEGORY = "MEDICINA"  # o título da aba: médicos
COLUMNS = ("uf", "periodo_ano", "periodo_mes", "categoria_profissional", "contagem")


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _period(value) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = _text(value)
    if isinstance(value, bool) or not re.fullmatch(r"\d{6}", text) or not 1 <= int(text[4:]) <= 12:
        raise ValueError("competência fora do formato AAAAMM: " + text[:20])
    return text


def _count(value, where: str) -> int:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, bool) or not isinstance(value, int):
        raise ValueError("contagem vazia ou não inteira em " + where)
    if value <= 0:
        raise ValueError("contagem não positiva em " + where)
    return value


def read_states(raw: bytes) -> dict:
    """Registros da aba Estados (uf, competência, contagem) na ordem da planilha, com as UFs e as notas da coluna A."""
    try:
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — qualquer falha de leitura recusa a planilha
        raise ValueError("a planilha não pôde ser lida como XLSX") from exc
    if SHEET not in book.sheetnames:
        raise ValueError('a planilha não tem a aba "' + SHEET + '"')
    rows = [list(row) for row in book[SHEET].iter_rows(values_only=True)]
    book.close()
    filled = lambda row, start=0: [_text(value) for value in row[start:] if _text(value)]  # noqa: E731
    if len(rows) < 2 or filled(rows[0]) != [TITLE] or _text(rows[0][0]) != TITLE:
        raise ValueError('o título da aba não é "' + TITLE + '"')
    if filled(rows[1]) != [SUBTITLE] or _text(rows[1][0]) != SUBTITLE:
        raise ValueError('o subtítulo da aba não é "' + SUBTITLE + '"')
    start = next((i for i, row in enumerate(rows) if row and _text(row[0]) == HEADER), None)
    if start is None or any(filled(row) for row in rows[2:start]):
        raise ValueError('a tabela precisa começar pelo cabeçalho "' + HEADER + '", logo abaixo do subtítulo')
    header = [_text(value) for value in rows[start][1:]]
    while header and not header[-1]:
        header.pop()
    siglas = sorted(code for code, _ in STATES.values())
    if sorted(header) != siglas:
        raise ValueError("o cabeçalho precisa trazer as 27 UFs, cada uma uma vez")
    records, seen, index = [], set(), start + 1
    while index < len(rows) and rows[index] and _text(rows[index][0]):
        row = rows[index]
        period = _period(row[0])
        if period in seen:
            raise ValueError("competência repetida: " + period)
        seen.add(period)
        if filled(row, 1 + len(header)):
            raise ValueError("célula preenchida fora da tabela na competência " + period)
        values = row[1:1 + len(header)] + [None] * max(0, len(header) + 1 - len(row))
        for uf, value in zip(header, values):
            records.append({"uf": uf, "competencia": period, "contagem": _count(value, uf + " " + period)})
        index += 1
    if not records:
        raise ValueError("a tabela não tem competências")
    notes = []
    for row in rows[index:]:
        if filled(row, 1):
            raise ValueError("célula preenchida fora da tabela, abaixo das competências")
        if row and _text(row[0]):
            notes.append(_text(row[0]))
    return {"registros": records, "ufs": header, "competencias": sorted(seen), "notas": notes}


def csv_bytes(table: dict) -> bytes:
    """O CSV da fonte estadual: competências em ordem cronológica e, em cada uma, as UFs na ordem da planilha."""
    order = {uf: position for position, uf in enumerate(table["ufs"])}
    lines = [";".join(COLUMNS)]
    for item in sorted(table["registros"], key=lambda item: (item["competencia"], order[item["uf"]])):
        period = item["competencia"]
        lines.append(";".join((item["uf"], period[:4], str(int(period[4:])), CATEGORY, str(item["contagem"]))))
    return ("\r\n".join(lines) + "\r\n").encode("utf-8")


def summary(raw: bytes, table: dict, content: bytes) -> dict:
    periods = table["competencias"]
    return {"aba": SHEET, "registros": len(table["registros"]), "ufs": len(table["ufs"]), "competencias": len(periods),
            "de": periods[0], "a": periods[-1], "notas": table["notas"],
            "sha256_planilha": hashlib.sha256(raw).hexdigest(), "arquivo": FILES["state"],
            "sha256_arquivo": hashlib.sha256(content).hexdigest(), "bytes_arquivo": len(content)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("planilha", type=Path)
    parser.add_argument("--gravar", action="store_true")
    args = parser.parse_args()
    planilha = args.planilha.read_bytes()
    tabela = read_states(planilha)
    conteudo = csv_bytes(tabela)
    print(json.dumps(summary(planilha, tabela, conteudo), ensure_ascii=False, indent=2))
    if args.gravar:
        (REPO_ROOT / FILES["state"]).write_bytes(conteudo)
        print("Fonte estadual gravada em " + FILES["state"] + "; regrave o manifesto, revise o diff e registre a mudança.")
