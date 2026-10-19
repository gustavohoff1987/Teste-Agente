"""Testes direcionados ao modelo XLSX e à leitura da hierarquia territorial."""
from io import BytesIO

import pandas as pd
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from main import app, read_dataset
from recommended_template import HEADERS, LEGEND_HEADER, recommended_workbook
from semantic_bridge import candidates, preflight, suggested_columns
from territorial_hierarchy import describe_hierarchy

client = TestClient(app)


def _filled_model(rename=False):
    wb = load_workbook(BytesIO(recommended_workbook()))
    ws = wb["Dados"]
    if rename:
        ws["A1"] = "Município"
        ws["B1"] = "UF"
        ws["C1"] = "Região Geográfica"
        ws["D1"] = "cod_ibge"
    row = ["Brasília", "DF", "Centro-Oeste", "000123", -15.8, -47.9,
           "Médicos", 2024, 6, 12, 5000]
    for col, value in enumerate(row, start=1):
        ws.cell(2, col, value)
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def test_download_xlsx_has_ordered_columns_and_right_hand_legend():
    response = client.get("/api/modelo/estrutura-recomendada")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    wb = load_workbook(BytesIO(response.content))
    assert wb.sheetnames == ["Dados"]
    ws = wb.active
    assert tuple(ws.cell(1, col).value for col in range(1, 12)) == HEADERS
    assert ws["M1"].value == LEGEND_HEADER
    assert ws["M3"].value == "territorio_analise"
    assert ws["D2"].number_format == "@"
    assert ws.freeze_panes == "A2"


def test_template_roundtrip_excludes_legend_and_preserves_code():
    frame, sheets = read_dataset(_filled_model(), "modelo.xlsx")
    assert sheets == ["Dados"]
    assert frame is not None
    assert len(frame) == 1
    assert all(column in frame.columns for column in HEADERS)
    assert "LEGENDA" not in str(frame.columns)
    assert str(frame.loc[0, "codigo_identificador"]) == "000123"
    assert frame.attrs["derived_fields"][0]["format"] == "yyyymm"
    mapping, profiles = candidates(frame)
    assert mapping["territory"] == "territorio_analise"
    # D1: o código é só sugestão até ser escolhido no seletor (seletor vazio com a sugestão no topo).
    assert "territory_id" not in mapping
    assert suggested_columns(profiles)["territory_id"] == ["codigo_identificador"]
    mapping["territory_id"] = "codigo_identificador"  # escolha do usuário
    hierarchy = preflight(frame, mapping, {"confirmed_roles": ["territory_id"]})["territorial_hierarchy"]
    assert [x["column"] for x in hierarchy["levels"]] == list(HEADERS[:3])
    assert hierarchy["aggregation_authorized"] is False
    assert hierarchy["relations"][0]["conflicting_units"] == 0


def test_renamed_brazilian_headers_preserve_generic_hierarchy():
    frame, _ = read_dataset(_filled_model(rename=True), "modelo.xlsx")
    assert frame is not None
    mapping, _ = candidates(frame)
    assert mapping["territory"] == "Município"
    assert "territory_id" not in mapping  # sugestão até ser escolhida
    mapping["territory_id"] = "cod_ibge"
    hierarchy = describe_hierarchy(frame, mapping)
    assert hierarchy["status"] == "detected"
    assert [x["column"] for x in hierarchy["levels"]] == [
        "Município", "UF", "Região Geográfica"
    ]


def test_homonymous_territory_names_with_distinct_ids_do_not_conflict():
    df = pd.DataFrame({
        "municipio": ["Santa Cruz", "Santa Cruz"],
        "uf": ["RN", "PB"],
        "regiao_geografica": ["Nordeste", "Nordeste"],
        "codigo_identificador": ["001", "002"],
    })
    mapping, _ = candidates(df)
    mapping["territory_id"] = "codigo_identificador"  # código escolhido no seletor (antes vinha pré-preenchido)
    hierarchy = describe_hierarchy(df, mapping)
    assert hierarchy["status"] == "detected"
    assert hierarchy["relations"][0]["conflicting_units"] == 0
