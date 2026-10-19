"""Conversão da planilha de médicos únicos por estado na fonte estadual (build_state_source.py).

As planilhas são sintéticas (valores sem significado real). Confere o formato gravado, que é o de
Dados_Estados_Brasil.csv, que nada é somado, estimado ou corrigido e que a planilha é recusada, sem gravar, quando a
estrutura que declara a contagem de médicos únicos por UF muda. O arquivo versionado é lido só para conferir que o
conversor o reproduz byte a byte.
"""
import hashlib
import io

import pytest
from openpyxl import Workbook

import build_state_source as builder
from brazil_data import FILES, REPO_ROOT, governed_manifest
from main import read_dataset
from territorial_selection import STATES

SIGLAS = [sigla for sigla, _ in STATES.values()]
PERIODS = ["202306", "202406"]


def workbook(*, sheet=builder.SHEET, title=builder.TITLE, subtitle=builder.SUBTITLE, header=None, rows=None,
             notes=("Fonte: planilha sintética de teste.",), change=None):
    book = Workbook()
    page = book.active
    page.title = sheet
    page.append([title])
    page.append([subtitle])
    page.append([])
    page.append([builder.HEADER] + (SIGLAS if header is None else header))
    for period, values in (rows if rows is not None else
                           [(int(period), [100 + i + 50 * n for i in range(27)]) for n, period in enumerate(PERIODS)]):
        page.append([period] + values)
    page.append([])
    for note in notes:
        page.append([note])
    if change:
        change(page)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_wide_sheet_becomes_one_row_per_state_and_competencia_without_changing_values():
    raw = workbook()
    table = builder.read_states(raw)
    assert (table["ufs"], table["competencias"], table["notas"]) == (SIGLAS, PERIODS, ["Fonte: planilha sintética de teste."])
    assert len(table["registros"]) == 54
    content = builder.csv_bytes(table)
    lines = content.decode("utf-8").split("\r\n")
    assert lines[0] == "uf;periodo_ano;periodo_mes;categoria_profissional;contagem" and lines[-1] == ""
    assert lines[1:3] == ["RO;2023;6;MEDICINA;100", "AC;2023;6;MEDICINA;101"]
    assert lines[28] == "RO;2024;6;MEDICINA;150" and len(lines) == 1 + 54 + 1
    # Cada célula da planilha vira uma linha, com o mesmo valor (nada somado nem estimado).
    expected = {(sigla, period): 100 + i + 50 * n for n, period in enumerate(PERIODS) for i, sigla in enumerate(SIGLAS)}
    frame, _ = read_dataset(content, FILES["state"])
    assert {(uf, period): int(value) for uf, period, value in
            zip(frame["uf"], frame["Período (ano e mês)"], frame["contagem"])} == expected
    assert set(frame["categoria_profissional"]) == {"MEDICINA"}
    report = builder.summary(raw, table, content)
    assert (report["registros"], report["ufs"], report["competencias"], report["de"], report["a"]) == (54, 27, 2, "202306",
                                                                                                         "202406")
    assert report["sha256_arquivo"] == hashlib.sha256(content).hexdigest() and report["arquivo"] == FILES["state"]
    with pytest.raises(ValueError, match="XLSX"):
        builder.read_states(b"planilha corrompida")


def test_versioned_state_source_is_reproduced_byte_by_byte_and_matches_the_manifest():
    raw = (REPO_ROOT / FILES["state"]).read_bytes()
    assert hashlib.sha256(raw).hexdigest() == governed_manifest()["sha256"]["state"]
    lines = raw.decode("utf-8").split("\r\n")
    assert lines[0] == ";".join(builder.COLUMNS) and lines[-1] == "" and "\n" not in "".join(lines)
    records = [line.split(";") for line in lines[1:-1]]
    assert len(records) == 594 and {record[3] for record in records} == {"MEDICINA"}
    assert all(int(record[4]) > 0 for record in records)
    periods = sorted({record[1] + record[2].zfill(2) for record in records})
    assert len(periods) == 22 and (periods[0], periods[-1]) == ("201401", "202406")
    ufs = list(dict.fromkeys(record[0] for record in records))
    assert sorted(ufs) == sorted(SIGLAS)
    table = {"ufs": ufs, "registros": [{"uf": record[0], "competencia": record[1] + record[2].zfill(2),
                                        "contagem": int(record[4])} for record in records]}
    assert builder.csv_bytes(table) == raw  # o mesmo formato, ordem e fim de linha do conversor


@pytest.mark.parametrize("arguments, message", [
    ({"sheet": "Municípios"}, 'não tem a aba "Estados"'),
    ({"title": "MÉDICOS POR ESTADO"}, "o título da aba"),
    ({"subtitle": "Soma dos municípios"}, "o subtítulo da aba"),
    ({"header": SIGLAS[:-1]}, "as 27 UFs"),
    ({"header": SIGLAS[:-1] + ["SP"]}, "as 27 UFs"),
    ({"rows": [(202313, [1] * 27)]}, "competência fora do formato AAAAMM: 202313"),
    ({"rows": [(202406, [1] * 27), (202406, [2] * 27)]}, "competência repetida: 202406"),
    ({"rows": [(202406, [1] * 26 + [None])]}, "contagem vazia ou não inteira em DF 202406"),
    ({"rows": [(202406, [1] * 26 + [1.5])]}, "contagem vazia ou não inteira em DF 202406"),
    ({"rows": [(202406, [1] * 26 + ["7"])]}, "contagem vazia ou não inteira em DF 202406"),
    ({"rows": [(202406, [0] + [1] * 26)]}, "contagem não positiva em RO 202406"),
    ({"rows": [(202406, [1] * 27 + [5])]}, "célula preenchida fora da tabela na competência 202406"),
    ({"rows": []}, "a tabela não tem competências"),
    ({"change": lambda page: page.cell(row=page.max_row, column=3, value="x")}, "abaixo das competências"),
    ({"change": lambda page: page.cell(row=3, column=1, value="nota fora do lugar")}, "logo abaixo do subtítulo"),
])
def test_changed_structure_is_refused_without_writing(arguments, message):
    before = hashlib.sha256((REPO_ROOT / FILES["state"]).read_bytes()).hexdigest()
    with pytest.raises(ValueError, match=message):
        builder.read_states(workbook(**arguments))
    assert hashlib.sha256((REPO_ROOT / FILES["state"]).read_bytes()).hexdigest() == before
