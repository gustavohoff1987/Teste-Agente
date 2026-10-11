"""Verificação integrada no navegador: UI → FastAPI → transporte Gemini simulado → validação → revisão → UI.

Cobre a conversa pelo backend real (A09–A13, troca de recorte nacional/municipal, falha e
recuperação do pré-exame, resposta concorrente), XLSX com escolha de aba, XLS, CSV malformado,
exportações e a minimização do que seria enviado ao provedor. Requer o catálogo Brasil instalado,
`npm run build` e, em outro terminal, um processo novo de:
    python backend/tests/serve_ui_fake_gemini.py
Depois, na raiz: python backend/tests/ui_integration_smoke.py
Não chama o Gemini: o transporte é substituído no servidor e a rede externa é bloqueada.
"""
import io
import json
import os
import sys
from pathlib import Path
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))
os.environ["GEMINI_API_KEY"] = ""  # somente neste processo; a chave local não é carregada
from openpyxl import Workbook  # noqa: E402
import xlwt  # noqa: E402
from playwright.sync_api import sync_playwright, expect  # noqa: E402
from ui_support import offline, stage, select  # noqa: E402
from llm_bridge import _sanitize  # noqa: E402

CALLS = Path(os.getenv("ADT_FAKE_GEMINI_LOG", ROOT / ".local" / "ui-validation" / "fake-gemini-calls.jsonl"))
URL = os.getenv("ADT_UI_URL", "http://127.0.0.1:18200/")
assert urlparse(URL).hostname in ("127.0.0.1", "localhost", "::1")
ROLE_STATES = {"recognized", "confirmed"}
CAP_STATES = {"available", "needs_confirmation", "unavailable"}
typed = []
results = []


def ok(label):
    results.append(label)
    print("PASS", label, flush=True)


def calls():
    if not CALLS.exists():
        return []
    return [json.loads(line) for line in CALLS.read_text(encoding="utf-8").splitlines() if line.strip()]


def field(request, name):
    body = request.post_data or ""
    marker = f'name="{name}"\r\n\r\n'
    if marker not in body:
        return None
    return body.split(marker, 1)[1].split("\r\n--", 1)[0]


def xlsx_two_sheets():
    wb = Workbook()
    notes = wb.active
    notes.title = "Notas"
    notes.append(["observacao"])
    notes.append(["aba sem dados analíticos"])
    data = wb.create_sheet("Dados")
    for row in (["municipio", "ano", "categoria", "quantidade"], ["Alfa", 2024, "Enfermagem", 12],
                ["Beta", 2024, "Enfermagem", 30]):
        data.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def xls_single():
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Plan1")
    for r, row in enumerate([["municipio", "ano", "categoria", "quantidade"], ["Alfa", 2024, "Enfermagem", 12],
                             ["Beta", 2024, "Enfermagem", 30]]):
        for c, value in enumerate(row):
            ws.write(r, c, value)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


with sync_playwright() as p:
    browser = p.chromium.launch(headless=True, channel=os.getenv("ADT_BROWSER_CHANNEL", "chrome"))
    ctx = browser.new_context(viewport={"width": 1440, "height": 1000}, accept_downloads=True)
    offline(ctx)
    page = ctx.new_page()
    page.set_default_timeout(45000)
    expect.set_options(timeout=45000)
    errors, posts = [], []
    page.on("pageerror", lambda e: errors.append(str(e)))
    page.on("request", lambda r: posts.append(r) if r.method == "POST" else None)
    log = page.get_by_role("log")
    prompt = page.get_by_role("textbox", name="Descreva sua análise ou faça uma pergunta")
    panel = page.get_by_role("region", name="Preparação da análise")

    def send(text, endpoint="interpret-context", status=200, ready=True):
        typed.append(text)
        if endpoint == "interpret-context" and status == 200 and ready:
            # Pedidos contextuais só são aceitos com pré-exame concluído (A10).
            expect(page.get_by_text("Base carregada: o Python confere", exact=False)).to_be_visible()
        prompt.fill(text)
        with page.expect_response(lambda r: r.url.endswith("/api/llm/" + endpoint)) as ev:
            page.get_by_role("button", name="Enviar mensagem").click()
        assert ev.value.status == status, (text, ev.value.status, ev.value.text()[:300])
        return ev.value

    def analyses():
        return [r for r in posts if r.url.endswith("/api/analyze")]

    page.goto(URL)
    page.get_by_role("button", name="Li e concordo", exact=True).click()

    # 1. Sem base: fluxo preliminar, sem verificação contextual.
    body = send("Quero analisar a densidade", endpoint="interpret").json()
    assert body["proposal"]["operation"] == "density" and "verification" not in body and body["executed"] is False
    expect(log).to_contain_text("Sugestão preliminar: Densidade")
    ok("sem base: proposta preliminar via /api/llm/interpret, sem verificação nem execução")

    # 2. Nova base reinicia o histórico.
    page.get_by_role("button", name="Usar dados do Brasil", exact=True).click()
    expect(page.get_by_text("Dados do Brasil carregados", exact=True)).to_be_visible()
    expect(log).not_to_contain_text("Sugestão preliminar: Densidade")
    expect(page.get_by_text("Sua conversa ·", exact=False)).to_have_count(0)
    ok("carregar base Brasil reinicia o histórico")

    # 3. Recorte nacional por Quantitativos e pedido contextual.
    select(page, "Quantitativos")
    expect(panel.get_by_text("Recorte atual: Brasil.", exact=False)).to_be_visible()
    first = "Quero quantitativos de MEDICINA em 202406"
    # 3a. A09 corrigido: período derivado ainda pendente não é tratado como valor ausente.
    response = send(first)
    request = response.request
    assert field(request, "dataset") == "brazil" and field(request, "scope") == "national"
    assert field(request, "context_ready") == "true" and 'name="file"' not in (request.post_data or "")
    v = response.json()["verification"]
    assert v["capability_status"] == "needs_confirmation", v
    assert any("Confirme a associação de Período" in b for b in v["blockers"]), v["blockers"]
    assert not any("correspondência única" in b for b in v["blockers"]), v["blockers"]
    expect(log).to_contain_text("O período mencionado não foi confirmado na base")
    ok("A09: período pendente → needs_confirmation com a pendência real, sem 'correspondência única'")
    # 3b. Após confirmar a associação do período, o mesmo pedido encontra os valores na base.
    stage(page, 1)
    panel.get_by_role("button", name="Confirmar associação", exact=True).click()
    expect(panel.get_by_role("button", name="Confirmar associação", exact=True)).to_have_count(0)
    stage(page, 3)
    expect(panel.get_by_text("Atualizando as pendências", exact=False)).to_have_count(0)
    stage(page, 1)
    response = send(first)
    body = response.json()
    v = body["verification"]
    assert v["source"] == "BACKEND_PREFLIGHT" and v["configuration_reviewed"] == "proposal"
    assert v["slots"]["period"] == "found_in_dataset" and v["slots"]["category"] == "found_in_dataset", v
    assert v["ready_to_execute"] is False and body["executed"] is False
    assert body["proposal"] == {"operation": "counts", "indicator": "quantity", "period": "202406",
                                "category": "MEDICINA", "territory": ""}, body["proposal"]
    expect(log).to_contain_text("Revisão determinística do Python")
    assert not analyses()
    ok(f"nacional: proposta revisada pelo backend (status {v['capability_status']}; atual {v['current_capability_status']})")

    # 4. A11: troca para o recorte municipal preserva o histórico, inclusive durante a inspeção.
    select(page, "Densidade")
    expect(log).to_contain_text(first)
    try:
        expect(page.get_by_text("Base em preparação", exact=False)).to_be_visible(timeout=3000)
        during = "aviso 'Base em preparação' exibido durante a inspeção"
    except AssertionError:
        during = "inspeção concluiu antes da observação do aviso"
    expect(panel.get_by_text("Recorte atual: Municípios.", exact=False)).to_be_visible()
    expect(panel.get_by_label("Unidade territorial de análise", exact=True)).to_have_value("municipio")
    expect(log).to_contain_text(first)
    expect(log).to_contain_text("Revisão determinística do Python")
    ok("A11: histórico preservado na troca nacional→municipal (" + during + ")")

    # 5. Fluxo 3 / A13: nova solicitação contextual no novo recorte, com retomada explícita.
    before = len(calls())
    second = "E a densidade no mesmo período?"
    response = send(second)
    request = response.request
    assert field(request, "scope") == "municipal" and field(request, "context_ready") == "true"
    assert json.loads(field(request, "history")) == [first, first]
    body = response.json()
    assert body["proposal"]["operation"] == "density" and body["proposal"]["period"] == "202406"
    assert body["proposal"]["category"] == "", body["proposal"]  # categoria sem retomada explícita
    v = body["verification"]
    assert v["configuration_reviewed"] == "proposal" and v["capability_status"] == "needs_confirmation", v
    assert not any("correspondência única" in b for b in v["blockers"]), v["blockers"]
    recorded = calls()[before:]
    assert len(recorded) == 1 and recorded[0]["messages"][-2:] == [_sanitize(first), _sanitize(second)]
    ok(f"fluxo 3: municipal com histórico; período retomado só por referência explícita (status {v['capability_status']})")

    response = send("Quero quantitativos")
    body = response.json()
    assert body["proposal"]["period"] == "" and body["proposal"]["category"] == "", body["proposal"]
    ok("A13: sem referência explícita, período/categoria anteriores não são herdados")

    # 6. Voltar ao nacional também preserva o histórico.
    select(page, "Quantitativos")
    expect(panel.get_by_text("Recorte atual: Brasil.", exact=False)).to_be_visible()
    expect(log).to_contain_text(second)
    expect(page.get_by_text("Sua conversa · 4 solicitações", exact=True)).to_be_visible()
    ok("A11: retorno municipal→nacional preserva as 4 solicitações")

    # 7. Fluxo 2: 'Ver requisitos' só abre a preparação; execução determinística exige confirmações.
    log.get_by_role("button", name="Ver requisitos de Densidade", exact=True).click()
    expect(panel.get_by_role("heading", name="Densidade", exact=True)).to_be_visible()
    expect(panel.get_by_text("Recorte atual: Municípios.", exact=False)).to_be_visible()
    stage(page, 3)
    expect(panel.get_by_role("button", name="Executar análise", exact=True)).to_be_disabled()
    assert not analyses()
    stage(page, 1)
    if panel.get_by_role("button", name="Confirmar associação", exact=True).count():
        panel.get_by_role("button", name="Confirmar associação", exact=True).click()
    if panel.get_by_role("button", name="Confirmar objeto de análise", exact=True).count():
        panel.get_by_role("button", name="Confirmar objeto de análise", exact=True).click()
    panel.get_by_role("button", name="Confirmar identificador territorial", exact=True).click()
    stage(page, 2)
    panel.get_by_label("O que o quantitativo representa?", exact=True).select_option("pessoas")
    panel.get_by_label("Confirmo que a população corresponde aos territórios e períodos analisados", exact=True).check()
    panel.get_by_label("Filtrar por período", exact=True).select_option("202406")
    stage(page, 3)
    execute = panel.get_by_role("button", name="Executar análise", exact=True)
    expect(execute).to_be_enabled()
    with page.expect_response(lambda r: r.url.endswith("/api/analyze")) as ev:
        execute.click()
    result = ev.value.json()
    assert ev.value.status == 200 and result["total_rows"] == 5570
    assert result["metadata"]["source"]["scope"] == "municipal"
    assert "Densidade de profissionais" not in page.locator('[role="log"]').text_content()
    ok("fluxo 2: proposta → preparação → execução determinística (5.570 municípios em 202406); chat sem resultados")
    page.get_by_role("button", name="Voltar à análise", exact=True).click()

    # 8. Fluxo 4 / A10: falha do pré-exame → 409 contextual, saudação local, recuperação.
    stage(page, 2)
    page.route("**/api/preflight", lambda route: route.fulfill(status=503, json={"detail": "Falha simulada"}))
    panel.get_by_label("Filtrar por período", exact=True).select_option("202401")
    expect(panel.get_by_role("button", name="Verificar novamente")).to_be_visible()
    before = len(calls())
    send("Quero a densidade em 202401", status=409)
    expect(page.get_by_role("alert").filter(has_text="Aguarde ou repita")).to_be_visible()
    send("Olá", ready=False)
    expect(log).to_contain_text("Que tipo de análise você deseja fazer?")
    assert len(calls()) == before
    page.unroute("**/api/preflight")
    panel.get_by_role("button", name="Verificar novamente").click()
    stage(page, 3)
    expect(panel.get_by_role("button", name="Executar análise", exact=True)).to_be_enabled()
    body = send("Quero a densidade em 202401").json()
    assert body["verification"]["source"] == "BACKEND_PREFLIGHT" and body["proposal"]["period"] == "202401"
    assert len(calls()) == before + 1
    ok("fluxo 4: falha → 409 sem chamar o provedor; saudação local; recuperação e novo pedido contextual")

    # 9. Concorrência: resposta pendente invalidada por mudança de parâmetros.
    turns = log.get_by_text("Proposta sujeita a confirmação").count()
    typed.append("lento quantitativos")
    prompt.fill("lento quantitativos")
    page.get_by_role("button", name="Enviar mensagem").click()
    stage(page, 2)
    panel.get_by_label("Filtrar por período", exact=True).select_option("202406")
    expect(page.get_by_role("button", name="Enviar mensagem")).to_be_enabled(timeout=15000)
    page.wait_for_timeout(500)
    assert log.get_by_text("Proposta sujeita a confirmação").count() == turns
    assert prompt.input_value() == "lento quantitativos"
    assert page.get_by_role("alert").filter(has_text="Não foi possível").count() == 0
    ok("concorrência: resposta obsoleta descartada, rascunho preservado, sem erro")

    # 10. A12: resposta incompatível do modelo é erro controlado, sem vazar credencial.
    response = send("invalido quantitativos", status=502)
    alert = page.get_by_role("alert").filter(has_text="fora do contrato")
    expect(alert).to_be_visible()
    assert "fake-offline-key" not in response.text() and "fake-offline-key" not in page.inner_text("body")
    assert log.get_by_text("Proposta sujeita a confirmação").count() == turns
    ok("A12: indicador incompatível → 502 controlado, sem turno novo e sem credencial exposta")

    # 11. Formatos: XLSX com duas abas, seleção de aba e reinício do histórico.
    page.get_by_role("button", name="Início", exact=True).click()
    page.locator('input[type="file"]').set_input_files({"name": "multi.xlsx", "buffer": xlsx_two_sheets(),
        "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
    sheet = page.get_by_label("Selecione a aba do arquivo Excel")
    expect(sheet).to_be_visible()
    expect(log).not_to_contain_text(first)
    send("Quero quantitativos", status=409)
    send("Olá", ready=False)
    sheet.select_option("Dados")
    expect(page.get_by_text("multi.xlsx carregado", exact=True)).to_be_visible()
    select(page, "Quantitativos")
    expect(panel.get_by_label("Unidade territorial de análise", exact=True)).to_have_value("municipio")
    expect(panel.get_by_label("Unidade temporal de análise", exact=True)).to_have_value("ano")
    ok("XLSX: escolha de aba exigida (pedido analítico 409 antes da aba), campos reconhecidos após a escolha")

    # 12. XLS de aba única: reconhecimento, execução, CSV e JSON de método.
    page.locator('input[type="file"]').set_input_files({"name": "plan.xls", "buffer": xls_single(),
        "mimeType": "application/vnd.ms-excel"})
    expect(page.get_by_text("plan.xls carregado", exact=True)).to_be_visible()
    expect(panel.get_by_label("Unidade territorial de análise", exact=True)).to_have_value("municipio")
    panel.get_by_role("button", name="Confirmar objeto de análise", exact=True).click()
    stage(page, 2)
    panel.get_by_label("O que o quantitativo representa?", exact=True).select_option("postos")
    stage(page, 3)
    with page.expect_response(lambda r: r.url.endswith("/api/analyze")) as ev:
        panel.get_by_role("button", name="Executar análise", exact=True).click()
    body = ev.value.json()
    assert ev.value.status == 200 and [r["Quantidade"] for r in body["rows"]] == [12, 30]
    with page.expect_download() as dl:
        page.get_by_role("button", name="CSV", exact=True).click()
    exported = Path(dl.value.path()).read_text(encoding="utf-8-sig").splitlines()
    assert len(exported) == 3 and exported[0].startswith('"Território"')
    with page.expect_download() as dl:
        page.get_by_role("button", name="Método e parâmetros", exact=True).click()
    method = json.loads(Path(dl.value.path()).read_text(encoding="utf-8"))
    assert method["operation"] == "counts" and method["options"]["count_unit"] == "postos"
    assert len(method["data_sha256"]) == 64 and any("postos" in n for n in method["notes"])
    page.get_by_role("button", name="Voltar à análise", exact=True).click()
    ok("XLS: reconhecimento, execução (12 e 30), exportação CSV e JSON de método")

    # 13. CSV estruturalmente inválido e extensão não suportada.
    page.get_by_role("button", name="Início", exact=True).click()
    page.locator('input[type="file"]').set_input_files({"name": "quebrado.csv", "mimeType": "text/csv",
        "buffer": b"Municipality;Year;Count\nA;2024;10;999\nB;2024;20\n"})
    expect(page.locator("#dados").get_by_role("alert").filter(has_text="CSV inconsistente na linha 2")).to_be_visible()
    expect(panel.get_by_role("alert").filter(has_text="CSV inconsistente na linha 2")).to_be_visible()
    expect(page.get_by_text("quebrado.csv carregado", exact=True)).to_have_count(0)
    expect(panel.get_by_role("button", name="Executar análise", exact=True)).to_have_count(0)
    send("Quero quantitativos", status=409)
    page.locator('input[type="file"]').set_input_files({"name": "dados.txt", "mimeType": "text/plain",
        "buffer": b"a;b\n1;2\n"})
    expect(page.locator("#dados").get_by_role("alert").filter(has_text="Formato não reconhecido")).to_be_visible()
    ok("CSV malformado rejeitado com linha indicada, sem execução; extensão não suportada recusada")

    # 14. Minimização: o provedor recebeu apenas textos digitados e estados enumerados.
    allowed = {_sanitize(t) for t in typed}
    for call in calls():
        for message in call["messages"]:
            if message.startswith("CONTEXTO_VERIFICADO_PELO_BACKEND: "):
                context = json.loads(message.split(": ", 1)[1])
                assert set(context) == {"source", "has_dataset", "recognized_roles", "capabilities"}, context
                assert set(context["recognized_roles"].values()) <= ROLE_STATES
                assert set(context["capabilities"].values()) <= CAP_STATES
            else:
                assert message in allowed, message
    ok(f"minimização: {len(calls())} chamadas ao transporte; só mensagens digitadas e estados enumerados")
    assert not errors, errors
    ok("nenhum erro JavaScript na página")
    browser.close()

print(f"INTEGRATION VERIFY PASS: {len(results)} verificações")
