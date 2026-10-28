"""Planilha fictícia na página única (28/10/2026, pedido do responsável): os nomes da planilha em todos os lugares.

Uma pessoa baixa o modelo de planilha pela janela de envio, preenche-o com territórios inventados (16 comunas agrupadas em 4
condados, com códigos, coordenadas no Atlântico Sul e cinco competências; os nomes foram conferidos em 28/10/2026 contra os
5.307 nomes de municípios e estados da base do Brasil, sem nenhuma coincidência), envia, dá os nomes no painel "Colunas"
(comuna/comunas, condado/condados) e percorre a página: a Seleção de Territórios, a Análise de Distribuição (Condições e
Exportar, com o CSV e o método), a Análise Espacial (a Escolha dos territórios, o resultado e "Entenda esta análise"), a
Análise de Tendência e a conversa. Onde a base do Brasil fala de município, estado e Brasil, a planilha fala de comuna,
condado e toda a planilha, com a concordância (comuna é feminino): "Todas que atendem", "Só as que eu marcar", "cada uma
ligada", "a referência mais próxima dela", "As demais"; o CSV tem a coluna "Condado"; a conversa reconhece o condado pelo
nome ("mostre só o condado Costa Ondina") e não promete um nível acima das comunas. Nenhum texto à vista traz município,
UF, IBGE, Brasil ou a sede municipal.

Catálogo sintético do Bloco 1 (a base do Brasil não é usada); servidor serve_ui_offline.py, sem o provedor de IA. A
planilha é criada numa pasta temporária a partir do modelo baixado pela tela.
"""
import csv
import io
import json
import os
import re
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright
from ui_support import AREA, URL, accept_privacy, is_post, offline, open_chat, options_of, send

OUTPUT = Path(__file__).resolve().parents[2] / ".local" / "ui-validation"
CONDADOS = {
    "Alto Vésper": ["Vesperina", "Rocaclara", "Brumavera", "Nevoral"],
    "Baixo Lúmen": ["Lumbral", "Trevalume", "Corvanha", "Pedrafina"],
    "Costa Ondina": ["Ondaflor", "Marolhado", "Salinésia", "Ventisca"],
    "Serra Quiral": ["Quiralto", "Solvento", "Cristalvo", "Fenduna"],
}
CENTROS = {"Alto Vésper": (-38.8, -20.6), "Baixo Lúmen": (-39.6, -19.4), "Costa Ondina": (-40.4, -18.2), "Serra Quiral": (-39.2, -17.6)}
PERIODOS = [(2022, 6), (2022, 12), (2023, 6), (2023, 12), (2024, 6)]
CABECALHO = ["territorio_analise", "territorio_superior_1", "territorio_superior_2", "codigo_identificador", "latitude",
             "longitude", "categoria_profissional", "periodo_ano", "periodo_mes", "quantidade", "populacao_estimada"]
# O que é próprio da organização brasileira e não pode aparecer com a planilha (os "três estados" da avaliação trivalente e o
# "estado atual do gráfico" não são territórios e ficam de fora).
DO_BRASIL = re.compile(r"munic[íi]pi|\bUFs?\b|\bIBGE\b|\bBrasil\b|sede municipal|"
                       r"\b(?:um|cada|o|do|no|nenhum) estado\b(?! (?:atual|salvo|de antes))|\btodos os estados\b|"
                       r"\bestados (?:da|do|de|escolhidos)\b|\bnível dos estados\b", re.IGNORECASE)
# Concordância com a unidade feminina: as formas de antes de 28/10/2026.
DISCORDANCIA = re.compile(r"\bos que não atendem\b|\bTodos que atendem\b|\bSó os que eu marcar\b|\bcada um ligado\b|\bdelea\b|"
                          r"\btodos os que atendem\b|\bOs demais\b|\btodos os comunas\b|\bos comunas\b|\bnível de condados\b")
TEXTOS = """() => {
  const textos = [document.body.innerText];
  for (const el of document.querySelectorAll('[aria-label],[title],[placeholder],[alt],option,optgroup,svg text,svg title')) {
    for (const nome of ['aria-label', 'title', 'placeholder', 'alt']) if (el.getAttribute(nome)) textos.push(el.getAttribute(nome));
    const tag = el.tagName.toLowerCase();
    if (tag === 'optgroup' && el.getAttribute('label')) textos.push(el.getAttribute('label'));
    if (tag === 'option' || el.namespaceURI === 'http://www.w3.org/2000/svg') textos.push(el.textContent || '');
  }
  return textos.join('\\n');
}"""


def preencher(modelo, destino):
    """O modelo baixado pela tela, com as colunas A:K preenchidas (a legenda à direita fica como veio)."""
    livro = load_workbook(modelo)
    folha = livro["Dados"]
    assert [folha.cell(row=1, column=i).value for i in range(1, 12)] == CABECALHO
    linha = 2
    for g, (condado, comunas) in enumerate(CONDADOS.items(), start=1):
        lat0, lon0 = CENTROS[condado]
        for u, comuna in enumerate(comunas, start=1):
            populacao = 5000 + 9000 * (g * 4 + u)
            for p, (ano, mes) in enumerate(PERIODOS):
                if comuna == "Nevoral" and ano < 2023:
                    continue  # sem registro no início da série
                for categoria, fator in (("Medicina", 0.4 + 0.25 * ((g + u) % 6)), ("Enfermagem", 1.5 + 0.3 * u)):
                    sem_coordenada = comuna == "Fenduna"
                    valores = [comuna, condado, "Zona Serrana" if g in (1, 4) else "Zona Marítima", f"SOL-{g:02d}{u:02d}",
                               None if sem_coordenada else round(lat0 + 0.1 * u, 5),
                               None if sem_coordenada else round(lon0 - 0.08 * u, 5), categoria, ano, mes,
                               int(populacao / 1000 * fator * (1 + 0.05 * p)), populacao + 100 * p]
                    for coluna, valor in enumerate(valores, start=1):
                        folha.cell(row=linha, column=coluna, value=valor)
                    linha += 1
    livro.save(destino)
    return destino


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="adt-planilha-ficticia-") as temporaria, sync_playwright() as p:
        pasta = Path(temporaria)
        browser = p.chromium.launch(headless=True, channel=os.getenv("ADT_BROWSER_CHANNEL", "chrome"))
        context = browser.new_context(viewport={"width": 1366, "height": 960}, accept_downloads=True)
        offline(context)
        page = context.new_page()
        falhas, vistos = [], {}
        page.on("pageerror", lambda erro: falhas.append(str(erro)))
        area = page.get_by_test_id(AREA["raiz"])
        chat = page.get_by_test_id("painel-chat")

        def aperta(locator):
            """Botões pelo teclado: confere o acesso sem mouse e evita a caixa da conversa por cima."""
            locator.scroll_into_view_if_needed()
            locator.focus()
            locator.press("Enter")

        pedidos = []
        page.on("request", lambda r: pedidos.append(r) if r.url.endswith("/api/painel") and r.method == "POST" else None)

        def painel(acao, condicao=None, timeout=180000):
            """A resposta de /api/painel do primeiro pedido feito depois da ação (com `condicao`, o primeiro cujas opções a
            atendem), como em ui_area_de_analise.py: antes, o cálculo em curso termina. A tela guarda as últimas respostas e só
            pede ao motor o que não está guardado."""
            page.wait_for_timeout(200)
            expect(area.get_by_test_id("analise-atualizando")).to_have_count(0, timeout=timeout)
            antes = len(pedidos)
            with page.expect_response(lambda r: is_post(r, "/api/painel") and any(r.request is pedido for pedido in pedidos[antes:])
                                      and (condicao is None or condicao(options_of(r.request))), timeout=timeout) as evento:
                acao()
            assert evento.value.status == 200, evento.value.text()[:2000]
            return evento.value.json(), options_of(evento.value.request)

        def guarda(nome):
            """Os textos à vista (e os atributos lidos por leitores de tela) para a varredura do fim."""
            page.wait_for_timeout(500)
            vistos[nome] = page.evaluate(TEXTOS)

        def baixa(locator):
            with page.expect_download(timeout=60000) as baixado:
                aperta(locator)
            destino = pasta / baixado.value.suggested_filename
            baixado.value.save_as(str(destino))
            return destino

        page.goto(URL)
        accept_privacy(page)
        # 1. O modelo baixado pela janela de envio, preenchido com os dados fictícios e enviado por ela.
        aperta(area.get_by_test_id("analise-fonte-arquivo"))
        janela = page.get_by_test_id("envio-de-dados")
        expect(janela).to_be_visible()
        modelo = baixa(janela.get_by_test_id("envio-modelo-baixar"))
        assert modelo.name == "modelo_territorial_recomendado.xlsx", modelo.name
        planilha = preencher(modelo, pasta / "solaria_comunas.xlsx")
        with page.expect_file_chooser() as escolha:
            janela.get_by_test_id("envio-escolher-arquivo").click()
        escolha.value.set_files(str(planilha))

        # 2. Colunas: o código e o grupo (sugeridos) escolhidos, os nomes, a unidade de contagem e o referencial.
        colunas = page.get_by_test_id("colunas-painel")
        expect(colunas).to_be_visible(timeout=90000)
        expect(colunas.get_by_test_id("colunas-territory")).to_have_value("territorio_analise")
        colunas.get_by_test_id("colunas-territory_id").select_option("codigo_identificador")
        colunas.get_by_test_id("colunas-grupo_territorial").select_option("territorio_superior_1")
        for campo, valor in (("unidade-singular", "comuna"), ("unidade-plural", "comunas"), ("grupo-singular", "condado"),
                             ("grupo-plural", "condados")):
            colunas.get_by_test_id("colunas-" + campo).fill(valor)
        colunas.get_by_test_id("colunas-unidade-de-contagem").select_option("pessoas")
        colunas.get_by_test_id("colunas-referencial").select_option("SIRGAS2000")
        expect(colunas.get_by_test_id("colunas-disponibilidade-espacial")).to_have_attribute("data-disponivel", "true", timeout=60000)
        colunas.get_by_test_id("colunas-comecar").click()
        expect(area.get_by_test_id("analise-aba-territorios")).to_be_visible(timeout=90000)
        expect(page.get_by_test_id("current-source")).to_contain_text("solaria_comunas.xlsx · comunas e condados")
        if chat.get_attribute("data-minimizada") != "true":
            chat.get_by_test_id("painel-chat-minimizar").click()

        # 3. Seleção de Territórios: os nomes da planilha, com a concordância; marcar todas calcula a Distribuição.
        expect(area.get_by_test_id("analise-busca")).to_have_attribute("placeholder", "Buscar comuna ou condado")
        expect(area.get_by_test_id("analise-marcar-todos")).to_have_text("Marcar todas as comunas (16)")
        resposta, _ = painel(lambda: aperta(area.get_by_test_id("analise-marcar-todos")))
        assert resposta["rotulos"]["unidade"] == {"singular": "comuna", "plural": "comunas", "genero": "f"}, resposta["rotulos"]
        expect(area.get_by_test_id("analise-aba-territorios-resumo")).to_have_text("Todas as comunas (16)")
        guarda("selecao")

        # 4. Distribuição: o título, as Condições e a exportação (a coluna do grupo no CSV e a nota das chaves no método).
        aperta(area.get_by_test_id("analise-familia-distribuicao"))
        expect(area.get_by_test_id("analise-titulo")).to_contain_text("Densidade de médicos nas comunas da planilha")
        expect(area.get_by_test_id("analise-subtitulo")).to_contain_text("Cada bolha é uma comuna.")
        aperta(area.get_by_test_id("analise-funcao-condicoes"))
        expect(area.get_by_test_id("analise-condicoes")).to_contain_text(
            "Separa as comunas que atendem a todas as condições escolhidas, as que não atendem e as que não têm dado para avaliar.")
        guarda("condicoes")
        aperta(area.get_by_test_id("analise-painel-fechar"))
        aperta(area.get_by_test_id("analise-funcao-exportar"))
        exportar = area.get_by_test_id("analise-painel")
        linhas = list(csv.reader(io.StringIO(baixa(exportar.get_by_test_id("analise-exportar-csv")).read_text(encoding="utf-8-sig")),
                                 delimiter=";"))
        assert linhas[0][:5] == ["Identificador", "Tipo", "Território", "Condado", "Competência"], linhas[0]
        assert "UF" not in linhas[0], linhas[0]
        vesperina = next(linha for linha in linhas[1:] if linha[2] == "Vesperina [SOL-0101]")
        assert vesperina[1:4] == ["comuna", "Vesperina [SOL-0101]", "Alto Vésper"], vesperina
        metodo = baixa(exportar.get_by_test_id("analise-exportar-metodo")).read_text(encoding="utf-8")
        assert ("Os nomes das chaves abaixo são os do contrato da API, iguais em toda fonte: nesta planilha, “municipios” são as "
                "comunas, e “estados”, os condados.") in metodo, metodo[-1500:]
        aperta(area.get_by_test_id("analise-painel-fechar"))
        guarda("distribuicao")

        # 5. Espacial: a Escolha dos territórios (botões, opções e frase), o cálculo conforme a base, o resultado e "Entenda".
        aperta(area.get_by_test_id("analise-familia-espacial"))
        janela = area.get_by_test_id("analise-painel").get_by_test_id("analise-referencia")
        expect(janela.get_by_test_id("territorios-base-todos")).to_have_text("Todas que atendem (16)", timeout=60000)
        expect(janela.get_by_test_id("territorios-base-marcados")).to_have_text("Só as que eu marcar")
        expect(janela.get_by_test_id("territorios-frase-referencia")).to_have_text(
            "a mais próxima de cada uma, entre as comunas de toda a planilha")
        guarda("janela-base")
        aperta(janela.get_by_test_id("territorios-aba-referencia"))
        menu = janela.get_by_test_id("territorios-referencia-partir-menu")
        assert "Mesmo condado de cada comuna da base" in menu.locator("option").all_text_contents()
        menu.select_option("mesma_uf_da_base")
        aperta(janela.get_by_test_id("territorios-referencia-mais-densidade"))
        linha = janela.get_by_test_id("territorios-referencia-filtro").first
        linha.get_by_test_id("territorios-referencia-filtro-operador").select_option(">")
        linha.get_by_test_id("territorios-referencia-filtro-valor").fill("1")
        expect(janela.get_by_test_id("territorios-frase-referencia")).to_have_text(
            "a mais próxima de cada uma, entre as comunas do mesmo condado com densidade acima de 1")
        expect(janela.get_by_test_id("territorios-referencia-usadas")).to_have_text(
            "Para cada comuna da base, só a candidata mais próxima fica como referência, entre as do mesmo condado.")
        guarda("janela-referencia")
        espacial, _ = painel(lambda: aperta(janela.get_by_test_id("territorios-calcular")), condicao=lambda o: (
            (o.get("consulta") or {}).get("referencia") or {}).get("partir_de", {}).get("tipo") == "mesma_uf_da_base")
        assert espacial["referencia"]["tipo"] == "consulta", espacial["referencia"]
        expect(janela.get_by_test_id("territorios-referencia-usadas")).to_contain_text("Ficaram como referência ")
        expect(janela.get_by_test_id("territorios-referencia-usadas")).to_contain_text(
            "cada uma a mais próxima de alguma comuna da base no mesmo condado.")
        guarda("janela-calculada")
        aperta(area.get_by_test_id("analise-painel-fechar"))
        expect(area.get_by_test_id("analise-par-base")).to_contain_text("Base: 16 · cada uma ligada à sua referência mais próxima")
        expect(area.get_by_test_id("analise-titulo")).to_contain_text("· Toda a planilha ·")
        aperta(area.get_by_test_id("analise-entenda-botao"))
        apresentacao = area.get_by_test_id("analise-apresentacao")
        expect(apresentacao).to_contain_text("entre cada comuna da base e a referência mais próxima dela.")
        expect(apresentacao).to_contain_text("a escolha de todas as que atendem ou só das marcadas.")
        guarda("espacial")
        aperta(area.get_by_test_id("analise-entenda-botao"))

        # 6. Tendência: o lugar do título numa planilha inteira.
        painel(lambda: aperta(area.get_by_test_id("analise-familia-tendencia")), condicao=lambda o: "tendencia" in o)
        expect(area.get_by_test_id("analise-titulo")).to_contain_text("Tendência da densidade de médicos · Toda a planilha ·")
        guarda("tendencia")

        # 7. A conversa: a orientação sem nível acima das comunas, o condado reconhecido pelo nome e "As demais".
        aperta(area.get_by_test_id("analise-familia-distribuicao"))
        expect(area.get_by_test_id("analise-titulo")).to_contain_text("Densidade de médicos nas comunas da planilha")
        open_chat(page)
        expect(chat.get_by_test_id("painel-chat-guia").locator("p").first).to_have_text(
            "Peça mudanças no gráfico citando comunas pelo nome: comparar, juntar, usar como referência ou mostrar as comunas de "
            "um condado.")
        send(page, "Quais comunas têm a maior densidade?")
        expect(chat).to_contain_text("As demais: ")
        with page.expect_request(lambda r: r.url.endswith("/api/painel") and r.method == "POST", timeout=120000) as pedido:
            send(page, "Mostre só o condado Costa Ondina")
        assert options_of(pedido.value)["base"]["municipios"]["estados"] == ["Costa Ondina"], options_of(pedido.value)["base"]
        expect(chat).to_contain_text("Mostrei a densidade nas 4 comunas do condado Costa Ondina.")
        send(page, "Mostre os condados")
        expect(chat).to_contain_text("Esta base não tem um nível acima das comunas no gráfico: cada bolha é uma comuna.")
        guarda("conversa")

        # 8. Nada do Brasil e nenhuma concordância de antes em tudo o que esteve à vista.
        for nome, texto in vistos.items():
            achados = sorted({m.group(0) for m in DO_BRASIL.finditer(texto)}) + sorted({m.group(0) for m in DISCORDANCIA.finditer(texto)})
            assert not achados, (nome, achados)
        page.screenshot(path=str(OUTPUT / "planilha-ficticia.png"), full_page=True)
        assert not falhas, falhas
        browser.close()
    print("Planilha fictícia conferida: comunas e condados em todos os lugares percorridos, sem município, UF, IBGE nem Brasil.")


if __name__ == "__main__":
    main()
