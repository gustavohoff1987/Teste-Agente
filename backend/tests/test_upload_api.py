"""Testa a importação HTTP real, inclusive formatos CSV, XLSX e XLS."""
import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from main import app


client = TestClient(app)
HEADERS = ("Territorio", "Ano", "Categoria", "Quantidade", "Populacao")
ROWS = [
    ("A", 2024, "Medicos", 100, 50000),
    ("B", 2024, "Medicos", 50, 25000),
]
MAPPING = {
    "territory": "Territorio",
    "period": "Ano",
    "category": "Categoria",
    "count": "Quantidade",
    "population": "Populacao",
}


def csv_bytes():
    values = (";".join(map(str, line)) for line in [HEADERS, *ROWS])
    return ("\n".join(values) + "\n").encode("utf-8")


def xlsx_bytes(extra_sheet=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"
    ws.append(HEADERS)
    for line in ROWS:
        ws.append(line)
    if extra_sheet:
        wb.create_sheet("Outra aba")
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def xls_bytes():
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Dados")
    for row_index, values in enumerate([HEADERS, *ROWS]):
        for column_index, value in enumerate(values):
            ws.write(row_index, column_index, value)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize("extension,factory", [
    ("csv", csv_bytes), ("xlsx", xlsx_bytes), ("xls", xls_bytes),
])
def test_file_inspect_and_real_density_analysis(extension, factory):
    contents = factory()
    filename = f"base.{extension}"
    inspected = client.post("/api/inspect", files={"file": (filename, contents)})
    assert inspected.status_code == 200, inspected.text
    profile = inspected.json()
    assert profile["row_count"] == 2
    assert profile["suggested_mapping"]["territory"] == "Territorio"
    assert profile["requires_sheet"] is False

    response = client.post("/api/analyze", files={"file": (filename, contents)}, data={
        "operation": "density", "mapping": json.dumps(MAPPING),
        "options": json.dumps({"count_unit": "pessoas", "population_denominator_confirmed": True})
    })
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["total_rows"] == 2
    assert all(row["Densidade por 1.000"] == 2.0 for row in data["rows"])


def test_multisheet_requires_explicit_choice():
    contents = xlsx_bytes(extra_sheet=True)
    initial = client.post("/api/inspect", files={"file": ("base.xlsx", contents)})
    assert initial.status_code == 200
    assert initial.json()["requires_sheet"] is True
    assert initial.json()["sheets"] == ["Dados", "Outra aba"]
    selected = client.post("/api/inspect", files={"file": ("base.xlsx", contents)}, data={"sheet": "Dados"})
    assert selected.status_code == 200
    assert selected.json()["requires_sheet"] is False
    assert selected.json()["row_count"] == 2


def test_error_message_is_visible_via_api_for_invalid_excel():
    response = client.post("/api/inspect", files={"file": ("broken.xlsx", b"not an excel file")})
    assert response.status_code == 422
    assert "Excel" in response.json()["detail"]


def test_unsupported_format_reported():
    response = client.post("/api/inspect", files={"file": ("base.txt", b"abc")})
    assert response.status_code == 422
    assert "CSV, XLS ou XLSX" in response.json()["detail"]


def test_structurally_inconsistent_csv_is_rejected_with_the_line():
    """Linha com campos a mais nunca é absorvida nem deslocada em silêncio (antes coberto só no navegador legado)."""
    for rows in (b"Municipality;Year;Count\nA;2024;10;999\nB;2024;20\n", b"Municipality;Year;Count\nA;2024;10\nB;2024\n"):
        response = client.post("/api/inspect", files={"file": ("quebrado.csv", rows)})
        assert response.status_code == 422, response.text
        line = 2 if b";999" in rows else 3
        assert f"CSV inconsistente na linha {line}" in response.json()["detail"]


def test_health_before_file_upload():
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"
