"""Planilha XLSX recomendada, baseada no esquema municipal e independente do país."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# Identificador legível, usado somente para excluir a legenda na importação.
LEGEND_HEADER = "LEGENDA DO MODELO — NÃO IMPORTAR"
HEADERS = (
    "territorio_analise", "territorio_superior_1", "territorio_superior_2",
    "codigo_identificador", "latitude", "longitude", "categoria_profissional",
    "periodo_ano", "periodo_mes", "quantidade", "populacao_estimada",
)
LEGEND = (
    ("territorio_analise", "Unidade territorial dos registros analisados.", "Município, distrito ou comunidade"),
    ("territorio_superior_1", "Território que contém o território de análise.", "UF, província ou departamento"),
    ("territorio_superior_2", "Território que contém o território superior 1.", "Região geográfica ou macrorregião"),
    ("codigo_identificador", "Código da unidade de análise; preserve zeros iniciais.", "Código local ou oficial"),
    ("latitude", "Latitude em graus decimais, de -90 a 90.", "Opcional fora de análises espaciais"),
    ("longitude", "Longitude em graus decimais, de -180 a 180.", "Opcional fora de análises espaciais"),
    ("categoria_profissional", "Profissão ou categoria observada na linha.", "Médicos, enfermeiros etc."),
    ("periodo_ano", "Ano da observação (quatro dígitos).", "2024"),
    ("periodo_mes", "Mês da observação, de 1 a 12.", "6"),
    ("quantidade", "Quantidade na unidade de contagem declarada.", "Número não negativo"),
    ("populacao_estimada", "Denominador populacional do território e período.", "Habitantes; exigido para densidade"),
)


def recommended_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    navy = "193451"
    pale = "EDF4FA"
    rule = Side(style="hair", color="CFDBE6")

    widths = (27, 28, 28, 27, 16, 16, 27, 16, 16, 18, 22)
    for index, (header, width) in enumerate(zip(HEADERS, widths), start=1):
        cell = ws.cell(row=1, column=index, value=header)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=10, color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=Side(style="medium", color="4CA6B5"))
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[1].height = 38
    ws.auto_filter.ref = "A1:K1001"

    # As células vazias são editáveis; não há linhas fictícias misturadas aos dados.
    for row in range(2, 102):
        for col in range(1, 12):
            c = ws.cell(row=row, column=col)
            c.fill = PatternFill("solid", fgColor="FFFFFF" if row % 2 == 0 else pale)
            c.border = Border(bottom=rule)
            c.alignment = Alignment(vertical="center")
            if col == 4:
                c.number_format = "@"  # identificadores opacos, inclusive zeros iniciais
            elif col in (5, 6):
                c.number_format = "0.000000"
            elif col in (8, 9, 10):
                c.number_format = "0"
            elif col == 11:
                c.number_format = "#,##0"
        ws.row_dimensions[row].height = 21

    rules = (
        ('H2:H10001', DataValidation(type="whole", operator="between", formula1=1000, formula2=9999,
                                     allow_blank=True)),
        ('I2:I10001', DataValidation(type="whole", operator="between", formula1=1, formula2=12,
                                     allow_blank=True)),
        ('E2:E10001', DataValidation(type="decimal", operator="between", formula1=-90, formula2=90,
                                     allow_blank=True)),
        ('F2:F10001', DataValidation(type="decimal", operator="between", formula1=-180, formula2=180,
                                     allow_blank=True)),
        ('J2:J10001', DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0,
                                     allow_blank=True)),
        ('K2:K10001', DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0,
                                     allow_blank=True)),
    )
    for target, validation in rules:
        validation.errorTitle = "Valor fora do intervalo"
        validation.error = "Confira a legenda ao lado da tabela."
        validation.showErrorMessage = True
        ws.add_data_validation(validation)
        validation.add(target)

    # À direita, no mesmo arquivo/aba, mas explicitamente fora da tabela A:K.
    ws.merge_cells("M1:O1")
    title = ws["M1"]
    title.value = LEGEND_HEADER
    title.fill = PatternFill("solid", fgColor=navy)
    title.font = Font(name="Aptos", color="FFFFFF", bold=True, size=11)
    title.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions["L"].width = 4
    ws.column_dimensions["M"].width = 28
    ws.column_dimensions["N"].width = 57
    ws.column_dimensions["O"].width = 35
    for index, heading in enumerate(("Campo", "Para que serve", "Exemplo ou orientação"), start=13):
        cell = ws.cell(2, index, heading)
        cell.font = Font(name="Aptos", bold=True, color=navy, size=10)
        cell.fill = PatternFill("solid", fgColor="DFEBF5")
        cell.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 25
    for row, (name, description, example) in enumerate(LEGEND, start=3):
        ws.cell(row, 13, name)
        ws.cell(row, 14, description)
        ws.cell(row, 15, example)
        for col in (13, 14, 15):
            cell = ws.cell(row, col)
            cell.font = Font(name="Aptos", size=10, color=navy if col == 13 else "32485D",
                             bold=col == 13)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.fill = PatternFill("solid", fgColor="F3F7FA" if row % 2 else "FFFFFF")
        ws.row_dimensions[row].height = 38
    ws.merge_cells("M16:O16")
    ws["M16"] = "Hierarquia: território de análise ⊂ superior 1 ⊂ superior 2."
    ws.merge_cells("M17:O17")
    ws["M17"] = "Os cabeçalhos podem ser adaptados; sinônimos conhecidos são reconhecidos. Outros exigem associação manual."
    ws.merge_cells("M18:O18")
    ws["M18"] = "Os níveis superiores são opcionais; sua presença NÃO autoriza somar contagens com sobreposição."
    ws.merge_cells("M19:O19")
    ws["M19"] = "Preencha somente A:K. A legenda M:O é ignorada na importação deste modelo."
    for row in range(16, 20):
        ws[f"M{row}"].font = Font(name="Aptos", italic=True, color=navy, size=10)
        ws[f"M{row}"].alignment = Alignment(wrap_text=True)
        ws.row_dimensions[row].height = 31
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
